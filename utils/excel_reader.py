from openpyxl import load_workbook

def read_excel_data(file_path, sheet_name):
    workbook = load_workbook(filename=file_path)
    sheet = workbook[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    headers = rows[0]
    data_rows = rows[1:]
    data_list = []
    for row in data_rows:
        data_list.append(dict(zip(headers, row)))
    return data_list
